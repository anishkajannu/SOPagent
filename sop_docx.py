"""Render a controlled SOP as a Word (.docx) document in Netramind house style.

Matches the real SOPs: NetraMind logo + bordered metadata header block on every
page, a bold Title line, numbered sections (1 Purpose … 6 Procedure) with hanging
indents, definition terms underlined, internal cross-references in red, a Revision
History table, and a Signature Page. Content comes in as the structured draft dict
the LLM produces; structure/numbering/formatting are fixed here.
"""

from __future__ import annotations

import datetime
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Calibri"
SIZE = 11
GREY = "D9D9D9"
BLACK = RGBColor(0x00, 0x00, 0x00)        # force true black (ignore theme text colour)
RED = RGBColor(0xC0, 0x00, 0x00)          # internal cross-reference color
LOGO_CANDIDATES = ["netramind_logo.png", "assets/netramind_logo.png"]

# A reference that points at another controlled document -> render in red.
_INTERNAL_REF = re.compile(r"^\s*(SOP|WIN|MAN|QM|POL|FRM|TMP|NM)\b", re.I)
# Strip a leading document number the model may have put in the title.
_TITLE_PREFIX = re.compile(r"^\s*(SOP|WIN|MAN|QM|POL)-\d{3}\s*[:\-–—]?\s*", re.I)


# ---------- low-level helpers ----------
def _font(run, size=SIZE, bold=False, italic=False, underline=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    # Default every run to true black so nothing inherits a grey theme text colour.
    run.font.color.rgb = color if color is not None else BLACK
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT)


def _add_run(paragraph, text, **kw):
    r = paragraph.add_run(text)
    _font(r, **kw)
    return r


def _page_field(paragraph):
    def field(instr):
        r = paragraph.add_run()
        _font(r)
        begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
        end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
        r._r.append(begin); r._r.append(it); r._r.append(end)

    field("PAGE")
    _add_run(paragraph, " of ")
    field("NUMPAGES")


def _set_table_borders(table, color="000000", sz=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)


def _shade(cell, fill=GREY):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _fixed_layout(table):
    """Lock column widths so Word honours them exactly (no auto-resize / clipping)."""
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def _zero_cell_margins(table):
    """Remove a table's internal cell padding (keeps a nested table from overflowing)."""
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "start", "bottom", "end", "left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), "0"); el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _cell_text(cell, text, bold=False, align=None, shade=False, size=SIZE):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    if align is not None:
        p.alignment = align
    _add_run(p, text, bold=bold, size=size)
    if shade:
        _shade(cell)


def _find_logo(base: Path):
    for name in LOGO_CANDIDATES:
        p = base / name
        if p.exists():
            return p
    return None


# ---------- header / footer ----------
def _build_header(section, doc_number, title, version, effective, logo):
    header = section.header
    header.is_linked_to_previous = False
    outer = header.add_table(rows=1, cols=2, width=Inches(6.4))
    outer.autofit = False
    _fixed_layout(outer)
    _zero_cell_margins(outer)          # so the nested meta table isn't pushed off the page
    logo_cell, meta_cell = outer.rows[0].cells
    logo_cell.width = Inches(2.6); meta_cell.width = Inches(3.8)
    logo_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER  # center against the table

    lp = logo_cell.paragraphs[0]
    if logo is not None:
        lp.add_run().add_picture(str(logo), width=Inches(2.4))  # fits inside the 2.6" cell
    else:
        _add_run(lp, "NetraMind", bold=True, size=16, color=RGBColor(0x1A, 0x56, 0xDB))

    meta = meta_cell.add_table(rows=5, cols=2)
    meta.autofit = False
    _fixed_layout(meta)
    for row in meta.rows:
        row.cells[0].width = Inches(1.5); row.cells[1].width = Inches(2.2)  # 3.7" total, slack inside 3.8"
    _set_table_borders(meta)
    meta.rows[0].cells[0].merge(meta.rows[0].cells[1])
    _cell_text(meta.rows[0].cells[0], "STANDARD OPERATING PROCEDURE",
               align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    rows = [("Document Number:", doc_number), ("Effective Date:", effective),
            ("Version Number:", version), ("Page:", None)]
    for i, (label, value) in enumerate(rows, start=1):
        _cell_text(meta.rows[i].cells[0], label, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if value is None:
            meta.rows[i].cells[1].text = ""
            _page_field(meta.rows[i].cells[1].paragraphs[0])
        else:
            _cell_text(meta.rows[i].cells[1], value)

    tp = header.add_paragraph()
    _add_run(tp, "Title: ", bold=True, color=BLACK)
    _add_run(tp, title, bold=True, color=BLACK)


def _build_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p1 = footer.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(p1, "THIS DOCUMENT CONTAINS CONFIDENTIAL AND PROPRIETARY INFORMATION", size=9)
    p2 = footer.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT   # page number on the right, like the official SOPs
    _add_run(p2, "Page ")                     # default black Calibri, same as the body
    _page_field(p2)


# ---------- body ----------
def _section_heading(doc, num, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.35))  # small gap: "1  PURPOSE"
    _add_run(p, str(num), bold=True)
    _add_run(p, "\t", bold=True)
    _add_run(p, title, bold=True)


def _body(doc, text):
    """A plain indented body paragraph (Purpose, Scope) aligned under the heading word."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(6)
    _add_run(p, text)


def _num_clause(doc, label, text, level=0):
    """A numbered clause (5.1, 6.1, …) with a hanging indent."""
    left = 0.9 + level * 0.4
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(left)
    p.paragraph_format.first_line_indent = Inches(-0.4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(left))  # text aligns with wrapped lines
    _add_run(p, f"{label}\t")
    _add_run(p, text)


def _definition(doc, label, text):
    """A definition line with the term underlined, like the real SOPs."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.9)
    p.paragraph_format.first_line_indent = Inches(-0.4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.9))
    _add_run(p, f"{label}\t")
    term, sep, rest = text.partition("—")   # em dash
    if not sep:
        term, sep, rest = text.partition(" - ")
    if sep:
        _add_run(p, term.strip(), underline=True)
        _add_run(p, " — ")
        _add_run(p, rest.strip())
    else:
        _add_run(p, text)


def _render_hier(doc, items, major):
    """Render a section (5=Responsibilities, 6=Procedure) with sub-subsections.

    `items` is already normalized by _as_hier_items: each entry is either a
    {'head', 'subs'} dict (renders as major.i + major.i.j) or a plain string.
    """
    for i, item in enumerate(items or ["[None]"], start=1):
        if isinstance(item, dict):
            _num_clause(doc, f"{major}.{i}", item.get("head", ""))
            for j, sub in enumerate(item.get("subs", []), start=1):
                _num_clause(doc, f"{major}.{i}.{j}", str(sub).strip(), level=1)
        else:
            _num_clause(doc, f"{major}.{i}", str(item).strip())


def _reference(doc, text):
    """A reference bullet; internal document references render in red."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.9)
    p.paragraph_format.first_line_indent = Inches(-0.4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.9))
    _add_run(p, "•\t")
    _add_run(p, text, color=RED if _INTERNAL_REF.match(text) else None)


def _clean_title(title: str) -> str:
    title = re.sub(r"^\s*title\s*:\s*", "", title, flags=re.I)
    title = _TITLE_PREFIX.sub("", title)
    return title.strip()


# ---------- input normalization (make rendering robust to model quirks) ----------
# The model's structured output SHOULD give lists, but it sometimes returns a single
# string (occasionally wrapped in <item>…</item> tags). Iterating a string yields one
# character per line — the "201-page" bug. These coerce any shape into clean lists.
_MAX_ITEMS = 60      # a real SOP section never has more than this; guards runaway output
_MAX_SUBS = 40


def _coerce_text(v) -> str:
    """One entry -> a single display string (handles str or a {term/meaning}-style dict)."""
    if isinstance(v, dict):
        term = v.get("term") or v.get("name") or v.get("role") or v.get("step") or ""
        meaning = (v.get("meaning") or v.get("definition") or v.get("description")
                   or v.get("text") or "")
        if term and meaning:
            return f"{term} — {meaning}"
        return (term or meaning or " ".join(str(x) for x in v.values())).strip()
    return str(v).strip()


def _as_str_list(value) -> list[str]:
    """Coerce definitions/references into a clean list of strings."""
    if value is None:
        return []
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return []
        tagged = re.findall(r"<item>(.*?)</item>", s, flags=re.S | re.I)
        if tagged:
            items = [t.strip() for t in tagged if t.strip()]
        else:
            s = re.sub(r"</?item>", "\n", s, flags=re.I)          # strip stray tags
            items = [re.sub(r"^[\s\-•*]+", "", p).strip() for p in re.split(r"[\n;]+", s)]
            items = [p for p in items if p] or [s]
        return items[:_MAX_ITEMS]
    if isinstance(value, list):
        out = [t for t in (_coerce_text(v) for v in value) if t]
        return out[:_MAX_ITEMS]
    return [str(value).strip()]


def _as_hier_items(value, head_key: str, sub_key: str) -> list:
    """Coerce responsibilities/procedure into a list of {head, subs} dicts or plain strings."""
    if value is None:
        return []
    if isinstance(value, str):
        return _as_str_list(value)          # a bare string -> flat string items (never chars)
    if isinstance(value, dict):
        value = [value]
    if not isinstance(value, list):
        return [str(value).strip()]
    out = []
    for item in value[:_MAX_ITEMS]:
        if isinstance(item, dict):
            subs_raw = item.get(sub_key) or item.get("duties") or item.get("substeps") or []
            if isinstance(subs_raw, str):
                subs = _as_str_list(subs_raw)
            elif isinstance(subs_raw, list):
                subs = [t for t in (_coerce_text(s) for s in subs_raw) if t][:_MAX_SUBS]
            else:
                subs = []
            head = (item.get(head_key) or item.get("role") or item.get("step") or "").strip()
            out.append({"head": head, "subs": subs})
        else:
            t = _coerce_text(item)
            if t:
                out.append(t)
    return out


def build_sop_docx(doc_number, draft, out_path, base_dir=".", version="01", effective="[On approval]"):
    base = Path(base_dir)
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(SIZE)

    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)
    sec.top_margin = Inches(1.6); sec.bottom_margin = Inches(0.9)

    title = _clean_title(draft.get("title", "").strip())
    _build_header(sec, doc_number, title, version, effective, _find_logo(base))
    _build_footer(sec)

    _section_heading(doc, 1, "PURPOSE")
    _body(doc, draft.get("purpose", "").strip())

    _section_heading(doc, 2, "SCOPE")
    _body(doc, draft.get("scope", "").strip())

    _section_heading(doc, 3, "REFERENCES")
    for ref in _as_str_list(draft.get("references")) or ["[None]"]:
        _reference(doc, ref)

    _section_heading(doc, 4, "DEFINITIONS")
    for i, d in enumerate(_as_str_list(draft.get("definitions")) or ["[None]"], start=1):
        _definition(doc, f"4.{i}", d)

    _section_heading(doc, 5, "RESPONSIBILITIES")
    _render_hier(doc, _as_hier_items(draft.get("responsibilities"), "role", "duties"), 5)

    _section_heading(doc, 6, "PROCEDURE")
    _render_hier(doc, _as_hier_items(draft.get("procedure"), "step", "substeps"), 6)

    rp = doc.add_paragraph(); rp.paragraph_format.space_before = Pt(12)
    _add_run(rp, "REVISION HISTORY", bold=True)
    rev = doc.add_table(rows=2, cols=3); _set_table_borders(rev)
    for c, h in zip(rev.rows[0].cells, ("Version Number", "Date", "Summary of Change(s)")):
        _cell_text(c, h, bold=True, shade=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    today = datetime.date.today().strftime("%d-%b-%Y").upper()   # e.g. 18-SEP-2026
    row = rev.rows[1].cells
    _cell_text(row[0], version, align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell_text(row[1], today, align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell_text(row[2], "Initial draft (pending QA approval)")

    doc.add_page_break()
    sp = doc.add_paragraph(); _add_run(sp, "SIGNATURE PAGE", bold=True)
    _add_run(doc.add_paragraph(), "[ADD ADDITIONAL ROWS AS NEEDED]")
    for role in ("Author", "Reviewer", "Approver"):
        rlp = doc.add_paragraph(); rlp.paragraph_format.space_before = Pt(8)
        _add_run(rlp, role + ":")
        t = doc.add_table(rows=2, cols=4); _set_table_borders(t)
        for c, h in zip(t.rows[0].cells, ("NAME", "TITLE", "SIGNATURE", "DATE")):
            _cell_text(c, h, bold=True, shade=True)
        for c in t.rows[1].cells:
            _cell_text(c, " ")

    doc.save(str(out_path))
    return str(out_path)
